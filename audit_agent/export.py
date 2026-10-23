from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import get_template
from .models import AuditReport


NAVY = RGBColor(20, 45, 73)
TEAL = RGBColor(16, 123, 120)
RISK_LABELS_VI = {
    "Low": "Thấp",
    "Moderate": "Trung bình",
    "High": "Cao",
    "Critical": "Nghiêm trọng",
}


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_labelled_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    paragraph.add_run(value)


def build_docx(report: AuditReport, payload: dict[str, Any]) -> bytes:
    """Create a polished Word draft suitable for a presentation download."""

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = NAVY

    draft = document.add_paragraph()
    draft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    draft_run = draft.add_run("DỰ THẢO (DRAFT) · CHỈ DÀNH CHO SOÁT XÉT NỘI BỘ")
    draft_run.bold = True
    draft_run.font.size = Pt(9)
    draft_run.font.color.rgb = TEAL

    title = document.add_paragraph(style="Title")
    title.add_run(report.report_title)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Kỳ kết thúc ngày {payload.get('period_end', 'Chưa cung cấp')}").italic = True

    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    engagement_rows = [
        ("Khách hàng", payload.get("client_name", "Chưa cung cấp")),
        (
            "Loại báo cáo",
            get_template(str(payload.get("report_type"))).get(
                "label_vi", payload.get("report_type", "Chưa cung cấp")
            ),
        ),
        ("Người lập", payload.get("prepared_by", "Chưa cung cấp")),
        ("Chủ nhiệm hợp đồng", payload.get("engagement_partner", "Chưa cung cấp")),
    ]
    for row, (label, value) in zip(table.rows, engagement_rows):
        row.cells[0].text = label
        row.cells[1].text = str(value or "Chưa cung cấp")
        _set_cell_shading(row.cells[0], "E8F1F1")
        row.cells[0].paragraphs[0].runs[0].bold = True

    document.add_heading("Tóm tắt điều hành", level=1)
    document.add_paragraph(report.executive_summary)
    document.add_heading("Mục tiêu", level=1)
    document.add_paragraph(report.objective)
    document.add_heading("Phạm vi", level=1)
    document.add_paragraph(report.scope)
    document.add_heading("Phương pháp thực hiện", level=1)
    document.add_paragraph(report.methodology)
    document.add_heading("Kết luận", level=1)
    document.add_paragraph(report.conclusion)

    document.add_heading("Phát hiện chi tiết", level=1)
    if not report.findings:
        document.add_paragraph("Chưa nhập phát hiện cần báo cáo.")
    for number, finding in enumerate(report.findings, start=1):
        document.add_heading(f"{number}. {finding.title}", level=2)
        _add_labelled_paragraph(document, "Lĩnh vực", finding.area)
        _add_labelled_paragraph(
            document, "Mức rủi ro", RISK_LABELS_VI.get(finding.risk_rating, finding.risk_rating)
        )
        _add_labelled_paragraph(document, "Hiện trạng", finding.condition)
        _add_labelled_paragraph(document, "Tiêu chí", finding.criteria)
        _add_labelled_paragraph(document, "Nguyên nhân", finding.cause)
        _add_labelled_paragraph(document, "Ảnh hưởng", finding.effect)
        _add_labelled_paragraph(document, "Kiến nghị", finding.recommendation)
        _add_labelled_paragraph(document, "Phản hồi của Ban Giám đốc", finding.management_response)
        _add_labelled_paragraph(document, "Tham chiếu bằng chứng", finding.evidence_reference)

    document.add_heading("Giới hạn", level=1)
    for item in report.limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Ghi chú soát xét", level=1)
    for item in report.reviewer_notes:
        document.add_paragraph(item, style="List Bullet")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Bản thử nghiệm Audit Automation Pipeline · Dự thảo phục vụ soát xét nội bộ có thẩm quyền")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
