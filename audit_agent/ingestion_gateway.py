from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


DEFAULT_GATEWAY_DB_PATH = Path("data/ingestion_gateway.sqlite3")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MIN_EXTRACTED_CHARACTERS = 80
MIN_FINE_TUNING_EXAMPLES = 10
ALLOWED_EXTENSIONS = {".docx", ".json", ".pdf", ".txt"}
USE_TYPES = ("RETRIEVAL", "EVALUATION", "TRAINING")

EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?84|0)[\s.-]?(?:\d[\s.-]?){8,10}(?!\d)")


@dataclass(frozen=True)
class ValidationMessage:
    level: Literal["warning", "blocker"]
    rule_code: str
    detail: str


@dataclass(frozen=True)
class UploadValidation:
    valid: bool
    clean_filename: str
    extension: str
    checksum: str
    extracted_text: str
    messages: list[ValidationMessage]


@dataclass(frozen=True)
class UploadReceipt:
    upload_id: str
    status: Literal["VALIDATED", "QUARANTINED"]
    checksum: str
    messages: list[ValidationMessage]


@dataclass(frozen=True)
class ExtractionPreview:
    """A non-persistent, reviewer-facing view of a received source file."""

    validation: UploadValidation
    character_count: int
    word_count: int
    line_count: int
    preview: str


@dataclass(frozen=True)
class DatasetExport:
    dataset_id: str
    use_type: str
    status: Literal["READY", "PREVIEW_ONLY"]
    item_count: int
    checksum: str
    filename: str
    content: bytes


def _connection(db_path: str | Path | None = None) -> sqlite3.Connection:
    path = Path(db_path or os.getenv("GATEWAY_DB_PATH", DEFAULT_GATEWAY_DB_PATH))
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def initialise_gateway(db_path: str | Path | None = None) -> None:
    with _connection(db_path) as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS gateway_upload (
                upload_id TEXT PRIMARY KEY,
                source_name TEXT NOT NULL,
                file_extension TEXT NOT NULL,
                media_type TEXT,
                byte_size INTEGER NOT NULL CHECK (byte_size >= 0),
                content_checksum TEXT NOT NULL UNIQUE,
                source_blob BLOB,
                extracted_text TEXT NOT NULL,
                report_type TEXT NOT NULL,
                training_input TEXT NOT NULL,
                anonymized INTEGER NOT NULL DEFAULT 0,
                uploaded_by TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                validation_status TEXT NOT NULL
                    CHECK (validation_status IN ('VALIDATED', 'QUARANTINED')),
                validation_json TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS gateway_use_decision (
                upload_id TEXT NOT NULL,
                use_type TEXT NOT NULL
                    CHECK (use_type IN ('RETRIEVAL', 'EVALUATION', 'TRAINING')),
                decision TEXT NOT NULL CHECK (decision IN ('APPROVED', 'REJECTED', 'REVOKED')),
                decided_by TEXT NOT NULL,
                decided_at TEXT NOT NULL,
                notes TEXT,
                PRIMARY KEY (upload_id, use_type),
                FOREIGN KEY (upload_id) REFERENCES gateway_upload(upload_id)
            );

            CREATE TABLE IF NOT EXISTS gateway_event (
                event_id INTEGER PRIMARY KEY AUTOINCREMENT,
                upload_id TEXT,
                event_status TEXT NOT NULL,
                rule_code TEXT,
                detail TEXT,
                occurred_at TEXT NOT NULL,
                FOREIGN KEY (upload_id) REFERENCES gateway_upload(upload_id)
            );

            CREATE TABLE IF NOT EXISTS gateway_dataset (
                dataset_id TEXT PRIMARY KEY,
                use_type TEXT NOT NULL
                    CHECK (use_type IN ('RETRIEVAL', 'EVALUATION', 'TRAINING')),
                dataset_status TEXT NOT NULL CHECK (dataset_status IN ('READY', 'PREVIEW_ONLY')),
                item_count INTEGER NOT NULL CHECK (item_count >= 0),
                content_checksum TEXT NOT NULL,
                created_by TEXT NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS gateway_dataset_item (
                dataset_id TEXT NOT NULL,
                upload_id TEXT NOT NULL,
                item_number INTEGER NOT NULL CHECK (item_number > 0),
                record_json TEXT NOT NULL,
                PRIMARY KEY (dataset_id, item_number),
                UNIQUE (dataset_id, upload_id),
                FOREIGN KEY (dataset_id) REFERENCES gateway_dataset(dataset_id),
                FOREIGN KEY (upload_id) REFERENCES gateway_upload(upload_id)
            );

            CREATE INDEX IF NOT EXISTS idx_gateway_upload_status_time
                ON gateway_upload(validation_status, uploaded_at DESC);
            CREATE INDEX IF NOT EXISTS idx_gateway_decision_use
                ON gateway_use_decision(use_type, decision);
            CREATE INDEX IF NOT EXISTS idx_gateway_dataset_time
                ON gateway_dataset(created_at DESC);
            """
        )


def _extract_docx(content: bytes) -> str:
    if not content.startswith(b"PK"):
        raise ValueError("The DOCX file does not have a valid Office document signature.")
    document = Document(BytesIO(content))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts)


def _extract_json(content: bytes) -> str:
    decoded = content.decode("utf-8-sig")
    parsed = json.loads(decoded)
    if not isinstance(parsed, (dict, list)):
        raise ValueError("The JSON root must be an object or array.")
    return json.dumps(parsed, ensure_ascii=False, indent=2)


def _extract_pdf(content: bytes) -> str:
    """Extract embedded text from a digitally generated PDF.

    This intentionally does not perform OCR. A scanned PDF without a text layer is
    rejected so it cannot be mistaken for a successfully extracted source.
    """
    if not content.startswith(b"%PDF-"):
        raise ValueError("The PDF file does not have a valid PDF signature.")
    try:
        reader = PdfReader(BytesIO(content))
    except PdfReadError as exc:
        raise ValueError(f"The PDF could not be read: {exc}") from exc
    if reader.is_encrypted:
        raise ValueError("Password-protected PDFs cannot be extracted in this prototype.")
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if not text:
        raise ValueError(
            "No embedded text was found. This may be a scanned PDF and needs OCR before intake."
        )
    return text


def _extract_text(extension: str, content: bytes) -> str:
    if extension == ".docx":
        return _extract_docx(content)
    if extension == ".json":
        return _extract_json(content)
    if extension == ".pdf":
        return _extract_pdf(content)
    if b"\x00" in content:
        raise ValueError("The text file contains binary null bytes.")
    return content.decode("utf-8-sig")


def validate_report_upload(filename: str, content: bytes) -> UploadValidation:
    clean_filename = Path(filename).name.strip()
    extension = Path(clean_filename).suffix.lower()
    checksum = hashlib.sha256(content).hexdigest()
    messages: list[ValidationMessage] = []
    extracted_text = ""

    if not clean_filename:
        messages.append(ValidationMessage("blocker", "FILENAME_REQUIRED", "A filename is required."))
    elif clean_filename != filename.strip():
        messages.append(
            ValidationMessage(
                "warning",
                "FILENAME_NORMALIZED",
                "Directory components were removed from the supplied filename.",
            )
        )
    if extension not in ALLOWED_EXTENSIONS:
        messages.append(
            ValidationMessage(
                "blocker",
                "FILE_TYPE_NOT_ALLOWED",
                "Only DOCX, digitally generated PDF, JSON, and UTF-8 TXT files are accepted in this prototype.",
            )
        )
    if not content:
        messages.append(ValidationMessage("blocker", "EMPTY_FILE", "The uploaded file is empty."))
    if len(content) > MAX_UPLOAD_BYTES:
        messages.append(
            ValidationMessage(
                "blocker",
                "FILE_TOO_LARGE",
                f"The file exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB prototype limit.",
            )
        )

    if not any(message.level == "blocker" for message in messages):
        try:
            extracted_text = _extract_text(extension, content).strip()
        except (UnicodeDecodeError, json.JSONDecodeError, ValueError, OSError) as exc:
            messages.append(ValidationMessage("blocker", "CONTENT_EXTRACTION_FAILED", str(exc)))

    if extracted_text and len(extracted_text) < MIN_EXTRACTED_CHARACTERS:
        messages.append(
            ValidationMessage(
                "blocker",
                "CONTENT_TOO_SHORT",
                f"Only {len(extracted_text)} characters were extracted; this does not look like a report.",
            )
        )
    if extracted_text and EMAIL_PATTERN.search(extracted_text):
        messages.append(
            ValidationMessage(
                "warning",
                "POTENTIAL_EMAIL",
                "The extracted report may contain an email address. Human anonymization review is required.",
            )
        )
    if extracted_text and PHONE_PATTERN.search(extracted_text):
        messages.append(
            ValidationMessage(
                "warning",
                "POTENTIAL_PHONE",
                "The extracted report may contain a phone number. Human anonymization review is required.",
            )
        )

    valid = not any(message.level == "blocker" for message in messages)
    return UploadValidation(valid, clean_filename, extension, checksum, extracted_text, messages)


def preview_report_upload(filename: str, content: bytes, *, preview_limit: int = 4_000) -> ExtractionPreview:
    """Extract and summarize an upload without writing its source bytes anywhere.

    This is deliberately separate from ``stage_report_upload`` so receiving staff can
    inspect what the extractor found before confirming authority to retain a source.
    """
    if preview_limit < 1:
        raise ValueError("The preview limit must be positive.")
    validation = validate_report_upload(filename, content)
    extracted = validation.extracted_text
    return ExtractionPreview(
        validation=validation,
        character_count=len(extracted),
        word_count=len(re.findall(r"\S+", extracted)),
        line_count=len(extracted.splitlines()) if extracted else 0,
        preview=extracted[:preview_limit],
    )


def stage_report_upload(
    filename: str,
    content: bytes,
    *,
    media_type: str | None,
    report_type: str,
    training_input: str,
    uploaded_by: str,
    upload_authorized: bool,
    db_path: str | Path | None = None,
) -> UploadReceipt:
    if not uploaded_by.strip():
        raise ValueError("The uploader identity is required.")
    if not upload_authorized:
        raise ValueError("The uploader must confirm authority to stage this report.")
    if not report_type.strip():
        raise ValueError("The report type is required.")

    validation = validate_report_upload(filename, content)
    upload_id = f"UPL-{uuid.uuid4().hex[:12].upper()}"
    status: Literal["VALIDATED", "QUARANTINED"] = (
        "VALIDATED" if validation.valid else "QUARANTINED"
    )
    now = datetime.now(timezone.utc).isoformat()
    messages_json = json.dumps(
        [message.__dict__ for message in validation.messages], ensure_ascii=False
    )
    initialise_gateway(db_path)

    try:
        with _connection(db_path) as connection:
            connection.execute(
                """
                INSERT INTO gateway_upload (
                    upload_id, source_name, file_extension, media_type, byte_size,
                    content_checksum, source_blob, extracted_text, report_type,
                    training_input, uploaded_by, uploaded_at, validation_status,
                    validation_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    upload_id,
                    validation.clean_filename or "unnamed",
                    validation.extension,
                    media_type,
                    len(content),
                    validation.checksum,
                    content if validation.valid else None,
                    validation.extracted_text,
                    report_type.strip(),
                    training_input.strip(),
                    uploaded_by.strip(),
                    now,
                    status,
                    messages_json,
                ),
            )
            connection.execute(
                """
                INSERT INTO gateway_event (
                    upload_id, event_status, rule_code, detail, occurred_at
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (
                    upload_id,
                    status,
                    "UPLOAD_VALIDATED" if validation.valid else "UPLOAD_QUARANTINED",
                    f"{len(validation.messages)} validation message(s)",
                    now,
                ),
            )
    except sqlite3.IntegrityError as exc:
        if "content_checksum" in str(exc):
            raise ValueError("This exact file has already been staged.") from exc
        raise

    return UploadReceipt(upload_id, status, validation.checksum, validation.messages)


def approve_gateway_use(
    upload_id: str,
    use_type: str,
    *,
    approved_by: str,
    anonymized: bool,
    authority_confirmed: bool,
    notes: str | None = None,
    db_path: str | Path | None = None,
) -> None:
    normalized_use = use_type.strip().upper()
    if normalized_use not in USE_TYPES:
        raise ValueError(f"Unknown use type: {use_type}")
    if not approved_by.strip():
        raise ValueError("An authorised approver is required.")
    if not authority_confirmed:
        raise ValueError("The approver must confirm authority for this AI use.")

    initialise_gateway(db_path)
    now = datetime.now(timezone.utc).isoformat()
    with _connection(db_path) as connection:
        upload = connection.execute(
            """
            SELECT validation_status, training_input
            FROM gateway_upload WHERE upload_id = ?
            """,
            (upload_id,),
        ).fetchone()
        if upload is None:
            raise ValueError("The selected upload does not exist.")
        if upload["validation_status"] != "VALIDATED":
            raise ValueError("A quarantined upload cannot be approved for AI use.")
        if normalized_use in {"EVALUATION", "TRAINING"} and not upload["training_input"].strip():
            raise ValueError(
                f"A realistic input/context example is required for {normalized_use.lower()} use."
            )
        if normalized_use == "TRAINING" and not anonymized:
            raise ValueError("Training approval requires an anonymized learning copy.")

        connection.execute(
            """
            INSERT INTO gateway_use_decision (
                upload_id, use_type, decision, decided_by, decided_at, notes
            ) VALUES (?, ?, 'APPROVED', ?, ?, ?)
            ON CONFLICT (upload_id, use_type) DO UPDATE SET
                decision = 'APPROVED',
                decided_by = excluded.decided_by,
                decided_at = excluded.decided_at,
                notes = excluded.notes
            """,
            (upload_id, normalized_use, approved_by.strip(), now, notes),
        )
        if anonymized:
            connection.execute(
                "UPDATE gateway_upload SET anonymized = 1 WHERE upload_id = ?",
                (upload_id,),
            )
        connection.execute(
            """
            INSERT INTO gateway_event (
                upload_id, event_status, rule_code, detail, occurred_at
            ) VALUES (?, 'APPROVED', ?, ?, ?)
            """,
            (
                upload_id,
                f"{normalized_use}_APPROVED",
                f"Approved by {approved_by.strip()}",
                now,
            ),
        )


def _dataset_record(row: sqlite3.Row, use_type: str) -> dict[str, Any]:
    metadata = {
        "upload_id": row["upload_id"],
        "source_name": row["source_name"],
        "report_type": row["report_type"],
        "content_checksum": row["content_checksum"],
        "anonymized": bool(row["anonymized"]),
    }
    if use_type == "RETRIEVAL":
        return {
            "document_id": row["upload_id"],
            "text": row["extracted_text"],
            "attributes": metadata,
        }
    if use_type == "EVALUATION":
        return {
            "item_id": row["upload_id"],
            "input": row["training_input"],
            "expected_output": row["extracted_text"],
            "metadata": metadata,
        }
    # Keep external training lines in the documented chat format. Lineage stays
    # in gateway_dataset_item rather than being added to the provider payload.
    return {
        "messages": [
            {"role": "user", "content": row["training_input"]},
            {"role": "assistant", "content": row["extracted_text"]},
        ]
    }


def build_dataset_export(
    use_type: str,
    *,
    created_by: str,
    db_path: str | Path | None = None,
) -> DatasetExport:
    normalized_use = use_type.strip().upper()
    if normalized_use not in USE_TYPES:
        raise ValueError(f"Unknown use type: {use_type}")
    if not created_by.strip():
        raise ValueError("The dataset creator is required.")

    initialise_gateway(db_path)
    with _connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT u.*
            FROM gateway_upload u
            JOIN gateway_use_decision d ON d.upload_id = u.upload_id
            WHERE u.validation_status = 'VALIDATED'
              AND d.use_type = ?
              AND d.decision = 'APPROVED'
              AND (? <> 'TRAINING' OR u.anonymized = 1)
            ORDER BY u.uploaded_at, u.upload_id
            """,
            (normalized_use, normalized_use),
        ).fetchall()
        if not rows:
            raise ValueError(f"No uploads are approved for {normalized_use.lower()} use.")

        records = [_dataset_record(row, normalized_use) for row in rows]
        content = (
            "\n".join(
                json.dumps(record, ensure_ascii=False, separators=(",", ":"))
                for record in records
            )
            + "\n"
        ).encode("utf-8")
        checksum = hashlib.sha256(content).hexdigest()
        dataset_id = f"DST-{uuid.uuid4().hex[:12].upper()}"
        status: Literal["READY", "PREVIEW_ONLY"] = (
            "PREVIEW_ONLY"
            if normalized_use == "TRAINING" and len(records) < MIN_FINE_TUNING_EXAMPLES
            else "READY"
        )
        now = datetime.now(timezone.utc).isoformat()
        connection.execute(
            """
            INSERT INTO gateway_dataset (
                dataset_id, use_type, dataset_status, item_count,
                content_checksum, created_by, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                dataset_id,
                normalized_use,
                status,
                len(records),
                checksum,
                created_by.strip(),
                now,
            ),
        )
        connection.executemany(
            """
            INSERT INTO gateway_dataset_item (
                dataset_id, upload_id, item_number, record_json
            ) VALUES (?, ?, ?, ?)
            """,
            [
                (
                    dataset_id,
                    row["upload_id"],
                    number,
                    json.dumps(record, ensure_ascii=False),
                )
                for number, (row, record) in enumerate(zip(rows, records), start=1)
            ],
        )
    return DatasetExport(
        dataset_id=dataset_id,
        use_type=normalized_use,
        status=status,
        item_count=len(records),
        checksum=checksum,
        filename=f"audit-pipeline-{normalized_use.lower()}-{dataset_id.lower()}.jsonl",
        content=content,
    )


def list_gateway_uploads(
    db_path: str | Path | None = None, *, limit: int = 100
) -> list[dict[str, Any]]:
    initialise_gateway(db_path)
    with _connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT
                u.upload_id,
                u.source_name,
                u.report_type,
                u.byte_size,
                u.content_checksum,
                u.validation_status,
                u.anonymized,
                u.uploaded_by,
                u.uploaded_at,
                COALESCE(GROUP_CONCAT(
                    CASE WHEN d.decision = 'APPROVED' THEN d.use_type END, ', '
                ), '') AS approved_uses
            FROM gateway_upload u
            LEFT JOIN gateway_use_decision d ON d.upload_id = u.upload_id
            GROUP BY u.upload_id
            ORDER BY u.uploaded_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
    return [dict(row) for row in rows]


def get_gateway_upload(
    upload_id: str, db_path: str | Path | None = None
) -> dict[str, Any] | None:
    initialise_gateway(db_path)
    with _connection(db_path) as connection:
        row = connection.execute(
            """
            SELECT upload_id, source_name, file_extension, media_type, byte_size,
                   content_checksum, report_type, training_input, anonymized,
                   uploaded_by, uploaded_at, validation_status, validation_json,
                   extracted_text, length(extracted_text) AS extracted_characters
            FROM gateway_upload WHERE upload_id = ?
            """,
            (upload_id,),
        ).fetchone()
    if row is None:
        return None
    result = dict(row)
    result["validation_messages"] = json.loads(result.pop("validation_json"))
    return result


def list_gateway_datasets(
    db_path: str | Path | None = None, *, limit: int = 50
) -> list[dict[str, Any]]:
    initialise_gateway(db_path)
    with _connection(db_path) as connection:
        rows = connection.execute(
            """
            SELECT dataset_id, use_type, dataset_status, item_count,
                   content_checksum, created_by, created_at
            FROM gateway_dataset
            ORDER BY created_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
    return [dict(row) for row in rows]
